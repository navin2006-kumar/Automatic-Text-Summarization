"""Document extraction and input validation for the summarization app."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path


SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx"}


class DocumentProcessingError(ValueError):
    """Raised when an uploaded document cannot be read or contains no text."""


def validate_text(text: str) -> str:
    """Return trimmed text, raising a useful error for empty input."""
    if not isinstance(text, str):
        raise DocumentProcessingError("Input must be text.")
    cleaned = text.strip()
    if not cleaned:
        raise DocumentProcessingError("Enter text or upload a document containing text.")
    return cleaned


def word_count(text: str) -> int:
    """Count words using whitespace tokenization for consistent UI statistics."""
    return len(text.split())


def extract_text(filename: str, content: bytes) -> str:
    """Extract text from a TXT, PDF, or DOCX upload held in memory."""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        allowed = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise DocumentProcessingError(
            f"Unsupported file type '{suffix or '(no extension)'}'. Use {allowed}."
        )
    if not content:
        raise DocumentProcessingError("The uploaded file is empty.")

    try:
        if suffix == ".txt":
            extracted = content.decode("utf-8-sig")
        elif suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(content))
            extracted = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            from docx import Document

            document = Document(BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            tables = [
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ]
            extracted = "\n".join(paragraphs + tables)
    except UnicodeDecodeError as exc:
        raise DocumentProcessingError("TXT files must use UTF-8 encoding.") from exc
    except Exception as exc:
        raise DocumentProcessingError(
            f"Could not read {suffix.upper()[1:]} document: {exc}"
        ) from exc

    return validate_text(extracted)
